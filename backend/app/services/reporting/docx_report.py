"""DOCX Compliance Report Generator (Owner: Aditya).

Builds editable .docx inspection reports with statutory metadata,
compliance verdicts, and declaration breakdown tables using python-docx.
"""

import io

from docx import Document
from docx.shared import Pt, RGBColor

from app.schemas.scan import ScanResponse


def generate_docx_report(scan: ScanResponse) -> bytes:
    """Generates an editable Word document audit report for the given scan."""
    doc = Document()

    # Document Title
    title = doc.add_heading("Legal Metrology Compliance Audit Report", level=0)
    title.paragraph_format.space_after = Pt(4)

    subtitle = doc.add_paragraph()
    sub_run = subtitle.add_run(f"Statutory Authority: {scan.rule_version} | Packaged Commodities Rules, 2011")
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph("―" * 55)

    # Overview Metadata Table
    doc.add_heading("Inspection Summary", level=1)
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.style = "Table Grid"

    cat_str = scan.category.value if scan.category is not None else "standard_retail"
    status_str = scan.overall_status.value if hasattr(scan.overall_status, "value") else str(scan.overall_status)

    rows_data = [
        ("Scan Identifier (UUID)", scan.scan_id),
        ("Product Name", scan.product_name or "N/A"),
        ("Product Category", cat_str),
        ("Audit Verdict", status_str),
        ("Confidence Score", f"{scan.overall_confidence * 100:.1f}%"),
    ]

    for i, (label, val) in enumerate(rows_data):
        meta_table.rows[i].cells[0].paragraphs[0].add_run(label).bold = True
        meta_table.rows[i].cells[1].paragraphs[0].add_run(val)

    doc.add_paragraph()

    # Compliance Findings Table
    doc.add_heading("Statutory Compliance Findings", level=1)
    results = scan.compliance_results or []
    comp_table = doc.add_table(rows=1 + len(results), cols=5)
    comp_table.style = "Table Grid"

    headers = ["Rule ID", "Rule Name", "Status", "Confidence", "Findings & Reason"]
    for col_idx, header in enumerate(headers):
        cell = comp_table.rows[0].cells[col_idx]
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.size = Pt(9)

    for row_idx, res in enumerate(results, start=1):
        row_cells = comp_table.rows[row_idx].cells
        row_cells[0].paragraphs[0].add_run(res.rule_id).font.size = Pt(8.5)
        row_cells[1].paragraphs[0].add_run(res.rule_name).font.size = Pt(8.5)

        st_run = row_cells[2].paragraphs[0].add_run(res.status.value if hasattr(res.status, "value") else str(res.status))
        st_run.font.size = Pt(8.5)
        st_run.bold = True

        row_cells[3].paragraphs[0].add_run(f"{res.confidence * 100:.0f}%").font.size = Pt(8.5)

        finding_text = res.violation_reason or res.measured_value or "Declaration verified"
        row_cells[4].paragraphs[0].add_run(finding_text).font.size = Pt(8.5)

    doc.add_paragraph()

    # Inspector Declaration Footer
    doc.add_heading("Inspector Declaration", level=1)
    p_decl = doc.add_paragraph(
        "This record was evaluated under statutory mandate of the Legal Metrology (Packaged Commodities) "
        "Rules, 2011. Results marked REVIEW_REQUIRED mandate physical examination using calibrated calipers "
        "and standard test weights."
    )
    p_decl.paragraph_format.space_after = Pt(12)

    # Save to buffer
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()
